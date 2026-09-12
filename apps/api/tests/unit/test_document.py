import io

from helpers_docx import build_fixture_docx

from rhapto.engine.document import (
    EDITABLE_ROLES,
    RawParagraph,
    apply_edits,
    classify,
    document_entities,
    document_numbers,
    number_key,
    parse_docx,
    quantity_tokens,
    section_kind_for,
    to_resume_document,
)
from rhapto.models.resume_document import ResumeHeader
from rhapto.models.source_document import DocParagraph, DocSection, Edit, SourceDocument


def _raw(
    text: str,
    style: str = "Normal",
    bold: bool = False,
    all_caps: bool = False,
    has_tab: bool = False,
    has_numbering: bool = False,
    centered: bool = False,
) -> RawParagraph:
    return RawParagraph(
        text=text,
        style=style,
        bold=bold,
        all_caps=all_caps,
        has_tab=has_tab,
        has_numbering=has_numbering,
        centered=centered,
    )


def test_parse_assigns_roles_and_sections() -> None:
    doc = parse_docx(build_fixture_docx(), "resume.docx")
    roles = [(p.role, p.text[:22]) for p in doc.paragraphs]
    assert roles[0] == ("name", "MAYA CHEN")
    assert roles[1][0] == "contact"
    assert ("heading", "PROFESSIONAL SUMMARY") in roles
    assert [r for r, _ in roles].count("summary") == 1
    assert [r for r, _ in roles].count("competency") == 1
    assert [r for r, _ in roles].count("entry_title") == 1
    assert [r for r, _ in roles].count("entry_org") == 1
    assert [r for r, _ in roles].count("bullet") == 2
    assert [r for r, _ in roles].count("skill") == 1
    assert [r for r, _ in roles].count("credential") == 1
    by_id = {p.id: p for p in doc.paragraphs}
    assert all(p.id.startswith("p") for p in doc.paragraphs)
    assert [s.heading for s in doc.sections] == [
        "PROFESSIONAL SUMMARY",
        "CORE COMPETENCIES",
        "PROFESSIONAL EXPERIENCE",
        "TECHNICAL SKILLS",
        "EDUCATION",
    ]
    summary = next(p for p in doc.paragraphs if p.role == "summary")
    assert summary.section == "PROFESSIONAL SUMMARY" and summary.id in doc.sections[0].paragraph_ids
    assert by_id[summary.id].text.startswith("Senior Data Program Manager")


def test_document_numbers_and_entities() -> None:
    doc = parse_docx(build_fixture_docx(), "resume.docx")
    numbers = document_numbers(doc)
    # Keys carry the unit class, so the document's "30%" does not license a bare "30" or a "$30M".
    assert {"8:plain", "12:plain", "30:percent", "2019:plain", "2025:plain"} <= numbers
    assert "30:plain" not in numbers and "12:currency" not in numbers
    # The contact line's phone digits are not quantities the resume claims.
    assert "555:plain" not in numbers and "0100:plain" not in numbers
    assert "Snowflake" in document_entities(doc) and "Acme Analytics" in document_entities(doc)


def test_number_key_classifies_units_and_scales() -> None:
    assert number_key("12") == "12:plain"
    assert number_key("30%") == "30:percent"
    assert number_key("30 percent") == "30:percent"
    assert number_key("8x") == "8:multiplier"
    assert number_key("$12M") == "12:currency"
    assert number_key("12M") == "12:scale"  # bare scale is not money
    assert number_key("50K") == "50:scale"
    assert number_key("$50K") == "50:currency"
    assert number_key("1,200") == "1200:plain"
    assert number_key("40 days") == "40:plain"


def test_quantity_tokens_normalises_ranges_and_spelled_scales() -> None:
    # A digit glued to an ASCII hyphen used to hide the right-hand number entirely.
    assert [key for _t, key in quantity_tokens("cost 30-85%")] == ["30:plain", "85:percent"]
    assert [key for _t, key in quantity_tokens("saved 12 million")] == ["12:scale"]
    assert [key for _t, key in quantity_tokens("saved $12 million")] == ["12:currency"]
    assert [key for _t, key in quantity_tokens("doubled the team")] == ["doubled"]


def test_apply_edits_and_mapping() -> None:
    doc = parse_docx(build_fixture_docx(), "resume.docx")
    bullet = next(p for p in doc.paragraphs if p.role == "bullet")
    edited = apply_edits(
        doc,
        [
            Edit(
                paragraph_id=bullet.id,
                before=bullet.text,
                after="Led the Snowflake migration across 12 teams.",
                reason="tighter",
            )
        ],
    )
    assert (
        next(p for p in edited.paragraphs if p.id == bullet.id).text
        == "Led the Snowflake migration across 12 teams."
    )
    assert doc.paragraphs != edited.paragraphs  # original untouched
    resume = to_resume_document(edited, ResumeHeader(name="Maya Chen"))
    assert (
        resume.summary[0].source_block_id
        == next(p for p in doc.paragraphs if p.role == "summary").id
    )
    exp = next(s for s in resume.sections if s.kind == "experience")
    assert (
        exp.entries[0].role == "Senior Data Program Manager"
        and exp.entries[0].period == "2019 – 2025"
    )
    assert (
        exp.entries[0].org == "Acme Analytics  |  Denver, CO" and len(exp.entries[0].bullets) == 2
    )
    assert exp.entries[0].bullets[0].source_block_id == bullet.id
    skills = next(s for s in resume.sections if s.kind == "skills")
    assert skills.entries[0].title in ("Program Delivery", "Platforms")
    assert EDITABLE_ROLES == frozenset({"summary", "competency", "skill", "bullet"})


def test_classify_empty_paragraph_is_other_with_stable_ids() -> None:
    raws = [
        _raw("Maya Chen", bold=True, all_caps=True, centered=True),
        _raw(""),
        _raw("SUMMARY", bold=True, all_caps=True),
    ]
    paragraphs = classify(raws)
    assert [p.id for p in paragraphs] == ["p0", "p1", "p2"]
    assert paragraphs[1].role == "other"
    assert paragraphs[1].text == ""


def test_classify_numbered_paragraph_is_bullet() -> None:
    raws = [
        _raw("Maya Chen", bold=True, all_caps=True, centered=True),
        _raw("EXPERIENCE", bold=True, all_caps=True),
        _raw("Shipped the migration.", has_numbering=True),
    ]
    paragraphs = classify(raws)
    assert paragraphs[2].role == "bullet"


def test_classify_heading_style_is_heading_even_without_bold_caps() -> None:
    raws = [
        _raw("Maya Chen"),
        _raw("Professional Summary", style="Heading 1"),
    ]
    paragraphs = classify(raws)
    assert paragraphs[1].role == "heading"
    assert paragraphs[1].text == "Professional Summary"


def test_classify_entry_title_without_tab() -> None:
    raws = [
        _raw("Maya Chen"),
        _raw("EXPERIENCE", bold=True, all_caps=True),
        _raw("Product Manager 2020 - 2022", bold=True),
    ]
    paragraphs = classify(raws)
    assert paragraphs[2].role == "entry_title"


def test_classify_multiline_contact_block() -> None:
    """Every paragraph between the name and the first heading that looks like contact info is
    tagged `contact`, not just the paragraph immediately after the name."""
    raws = [
        _raw("Maya Chen", bold=True, all_caps=True, centered=True),
        _raw("555 0100"),
        _raw("maya.chen@example.com"),
        _raw("SUMMARY", bold=True, all_caps=True),
    ]
    paragraphs = classify(raws)
    assert [p.role for p in paragraphs[1:3]] == ["contact", "contact"]
    assert paragraphs[3].role == "heading"


def test_entry_title_without_tab_splits_role_and_period() -> None:
    paragraphs = [
        DocParagraph(id="p0", text="Maya Chen", role="name", section=None),
        DocParagraph(id="p1", text="EXPERIENCE", role="heading", section="EXPERIENCE"),
        DocParagraph(
            id="p2", text="Product Manager 2020 - 2022", role="entry_title", section="EXPERIENCE"
        ),
    ]
    doc = SourceDocument(
        filename="r.docx",
        paragraphs=paragraphs,
        sections=[DocSection(heading="EXPERIENCE", paragraph_ids=["p2"])],
    )
    resume = to_resume_document(doc, ResumeHeader(name="Maya Chen"))
    exp = next(s for s in resume.sections if s.kind == "experience")
    assert exp.entries[0].role == "Product Manager"
    assert exp.entries[0].period == "2020 - 2022"


def test_to_resume_document_routes_bulleted_skills_and_credentials_by_section() -> None:
    """A section formatted with real Word bullets (not `Label: value` lines) still lands in the
    right ResumeSection, grouped one entry per section heading."""
    paragraphs = [
        DocParagraph(id="p0", text="Maya Chen", role="name", section=None),
        DocParagraph(id="p1", text="TECHNICAL SKILLS", role="heading", section="TECHNICAL SKILLS"),
        DocParagraph(id="p2", text="Snowflake", role="bullet", section="TECHNICAL SKILLS"),
        DocParagraph(id="p3", text="dbt", role="bullet", section="TECHNICAL SKILLS"),
        DocParagraph(id="p4", text="EDUCATION", role="heading", section="EDUCATION"),
        DocParagraph(id="p5", text="B.S. Computer Science", role="bullet", section="EDUCATION"),
    ]
    doc = SourceDocument(
        filename="r.docx",
        paragraphs=paragraphs,
        sections=[
            DocSection(heading="TECHNICAL SKILLS", paragraph_ids=["p2", "p3"]),
            DocSection(heading="EDUCATION", paragraph_ids=["p5"]),
        ],
    )
    assert section_kind_for(doc, paragraphs[2]) == "skill"
    assert section_kind_for(doc, paragraphs[5]) == "credential"
    resume = to_resume_document(doc, ResumeHeader(name="Maya Chen"))
    assert not any(s.kind == "experience" for s in resume.sections)
    skills = next(s for s in resume.sections if s.kind == "skills")
    assert len(skills.entries) == 1
    assert skills.entries[0].title == "TECHNICAL SKILLS"
    assert [b.text for b in skills.entries[0].bullets] == ["Snowflake", "dbt"]
    credentials = next(s for s in resume.sections if s.kind == "credentials")
    assert len(credentials.entries) == 1
    assert credentials.entries[0].title == "EDUCATION"
    assert [b.text for b in credentials.entries[0].bullets] == ["B.S. Computer Science"]


def test_bold_caps_contact_line_before_first_heading_is_contact() -> None:
    raws = [
        RawParagraph("MAYA CHEN", "Normal", True, True, False, False, True),
        RawParagraph("MAYA.CHEN@EXAMPLE.COM | 555 0100", "Normal", True, True, False, False, True),
        RawParagraph("PROFESSIONAL SUMMARY", "Normal", True, True, False, False, False),
    ]
    roles = [p.role for p in classify(raws)]
    assert roles == ["name", "contact", "heading"]


def test_raw_detects_word_numbering() -> None:
    """A Normal-style paragraph with a w:numPr element (Word's bullet/numbering) is a bullet."""
    from docx import Document
    from docx.oxml.ns import qn

    d = Document()
    d.add_paragraph("MAYA CHEN")
    numbered = d.add_paragraph("Shipped the thing")
    ppr = numbered._p.get_or_add_pPr()
    ppr.append(ppr.makeelement(qn("w:numPr"), {}))
    buf = io.BytesIO()
    d.save(buf)
    doc = parse_docx(buf.getvalue(), "r.docx")
    assert doc.paragraphs[1].role == "bullet"
