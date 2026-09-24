from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from helpers import bullet, demo_resume

from rhapto.engine.render.docx import OrphanBulletError, render_docx
from rhapto.profile.loader import load_profile


def test_renders_single_column_ats_safe_docx(demo_profile_dir: Path) -> None:
    blocks = load_profile(demo_profile_dir).block_map()
    data = render_docx(demo_resume(), blocks)
    doc = Document(BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "Maya Chen"
    assert "maya.chen@example.com | Denver, CO" in texts
    # Headings an ATS can classify. "CREDENTIALS" used to be here and carries no token any parser
    # looks for, so the block under it went unindexed; see test_render_templates.py for the rule.
    assert (
        "SUMMARY" in texts
        and "EXPERIENCE" in texts
        and "PROJECTS" in texts
        and "EDUCATION & CERTIFICATIONS" in texts
    )
    assert "Senior Data Program Manager | Acme Analytics | 2019-2025" in texts
    assert any("cutting warehouse cost 18%" in t for t in texts)
    assert len(doc.tables) == 0 and len(doc.inline_shapes) == 0
    bullet_styles = {
        p.style.name for p in doc.paragraphs if p.text.startswith("Led cross-functional")
    }
    assert bullet_styles == {"List Bullet"}
    assert doc.styles["Normal"].font.name == "Calibri"


def test_section_heading_comes_from_kind_not_llm_title(demo_profile_dir: Path) -> None:
    blocks = load_profile(demo_profile_dir).block_map()
    resume = demo_resume()
    resume.sections[0].title = "Relevant Professional Experience"
    doc = Document(BytesIO(render_docx(resume, blocks)))
    texts = [p.text for p in doc.paragraphs]
    assert "EXPERIENCE" in texts
    assert "RELEVANT PROFESSIONAL EXPERIENCE" not in texts


def test_orphan_bullet_raises(demo_profile_dir: Path) -> None:
    blocks = load_profile(demo_profile_dir).block_map()
    resume = demo_resume()
    resume.sections[0].entries[0].bullets.append(bullet("Made up.", "ghost"))
    with pytest.raises(OrphanBulletError) as exc:
        render_docx(resume, blocks)
    assert exc.value.block_id == "ghost" and exc.value.path == "sections[0].entries[0].bullets[2]"


def test_orphan_entry_raises(demo_profile_dir: Path) -> None:
    blocks = load_profile(demo_profile_dir).block_map()
    resume = demo_resume()
    resume.sections[1].entries[0].source_block_id = "ghost"
    with pytest.raises(OrphanBulletError):
        render_docx(resume, blocks)
