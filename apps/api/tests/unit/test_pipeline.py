import io
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from helpers import bullet, demo_extract, demo_resume
from helpers_docx import build_fixture_docx
from pydantic import ValidationError

from rhapto.engine.compose import AnswerItem, ComposeOutput
from rhapto.engine.document import parse_docx
from rhapto.engine.pipeline import CallBudget, LLMBudgetExceeded, TailorResult, tailor
from rhapto.engine.providers.fake import FakeEmbeddingProvider, FakeLLMProvider
from rhapto.engine.providers.llm import MalformedOutputError, TokenUsage
from rhapto.engine.tune import ProposedEdit, TuneOutput
from rhapto.engine.types import Profile, ProfileError, TailorRequest
from rhapto.models.source_document import SourceDocument
from rhapto.profile.loader import load_profile

JD = "ExampleCo seeks a Data Platform Program Manager to lead our Snowflake migration."


@pytest.fixture
def profile(demo_profile_dir: Path) -> Profile:
    return load_profile(demo_profile_dir)


def good_output() -> dict[str, Any]:
    resume = demo_resume()
    return ComposeOutput(
        summary=resume.summary,
        sections=resume.sections,
        cover_note="Dear ExampleCo team, " + "word " * 120,
        change_log="Emphasised migration.",
        answers=[AnswerItem(key="why_this_company", value="Data.")],
    ).model_dump(mode="json")


def bad_output() -> dict[str, Any]:
    output = good_output()
    output["sections"][0]["entries"][0]["bullets"][1] = bullet(
        "Cut warehouse cost 25%.", "acme-migration"
    ).model_dump()
    return output


def cover_note_metric_output() -> dict[str, Any]:
    output = good_output()
    output["cover_note"] = "I cut costs 37% for the platform team. " + "word " * 120
    return output


def missing_role_output() -> dict[str, Any]:
    """Drops the acme-data-pm entry entirely -- the adversarial case this project targets."""
    output = good_output()
    output["sections"][0]["entries"] = []
    return output


def repair_drops_credential_output() -> dict[str, Any]:
    """A 'repair' that fixes the role but drops the credential instead -- must still block."""
    output = good_output()
    output["sections"][2]["entries"] = []
    return output


def test_call_budget() -> None:
    budget = CallBudget(max_calls=2)
    budget.before_call()
    budget.after_call(TokenUsage(input_tokens=3))
    budget.before_call()
    budget.after_call(TokenUsage(input_tokens=4))
    assert budget.calls == 2 and budget.usage.input_tokens == 7
    with pytest.raises(LLMBudgetExceeded):
        budget.before_call()


async def test_happy_path_uses_two_calls(profile: Profile) -> None:
    llm = FakeLLMProvider([demo_extract(), good_output()])
    steps: list[str] = []

    async def on_step(name: str) -> None:
        steps.append(name)

    result = await tailor(
        TailorRequest(jd_text=JD), profile, llm, FakeEmbeddingProvider(), on_step=on_step
    )
    assert isinstance(result, TailorResult)
    package = result.package
    assert package.status == "draft" and package.guardrail_report.passed and package.llm_calls == 2
    assert package.version == 1 and package.track_id == "data-pm"
    assert package.job.company == "ExampleCo" and package.job.jd_text == JD
    assert package.resume.header.name == "Maya Chen"
    assert "acme-migration" in result.selection.block_ids
    assert result.docx[:2] == b"PK"
    assert steps == ["extract", "select", "compose", "validate", "render"]


MALFORMED = {"$PARAMETER_NAME": "$PARAMETER_VALUE"}  # seen verbatim from a forced tool call


async def test_malformed_compose_is_retried_within_the_budget(profile: Profile) -> None:
    llm = FakeLLMProvider([demo_extract(), MALFORMED, good_output()])
    result = await tailor(TailorRequest(jd_text=JD), profile, llm, FakeEmbeddingProvider())
    assert result.package.status == "draft" and result.package.llm_calls == 3
    assert len(llm.calls) == 3
    # I-3: three calls and a draft, but the third was a malformed-output retry, not a repair.
    assert result.pre_repair_report is None and result.repaired is False


async def test_malformed_compose_twice_raises_instead_of_a_fourth_call(profile: Profile) -> None:
    llm = FakeLLMProvider([demo_extract(), MALFORMED, MALFORMED, good_output()])
    with pytest.raises(MalformedOutputError):
        await tailor(TailorRequest(jd_text=JD), profile, llm, FakeEmbeddingProvider())
    assert len(llm.calls) == 3


async def test_malformed_repair_keeps_the_blocked_draft(profile: Profile) -> None:
    llm = FakeLLMProvider([demo_extract(), bad_output(), MALFORMED])
    result = await tailor(TailorRequest(jd_text=JD), profile, llm, FakeEmbeddingProvider())
    assert result.package.status == "blocked" and result.package.llm_calls == 3
    assert not result.package.guardrail_report.passed


async def test_repair_path_uses_three_calls_and_passes(profile: Profile) -> None:
    llm = FakeLLMProvider([demo_extract(), bad_output(), good_output()])
    result = await tailor(TailorRequest(jd_text=JD), profile, llm, FakeEmbeddingProvider())
    assert result.package.status == "draft" and result.package.llm_calls == 3
    assert result.repaired is True
    assert result.pre_repair_report is not None and not result.pre_repair_report.passed
    repair_call = llm.calls[2]
    assert "25%" in repair_call.messages[0].content and repair_call.output_schema is ComposeOutput
    assert repair_call.system == llm.calls[1].system  # same cached system blocks as compose


async def test_usage_reaches_the_package(profile: Profile) -> None:
    """Every FakeLLMProvider call reports input=10/output=5; two calls should sum onto the package."""
    llm = FakeLLMProvider([demo_extract(), good_output()])
    result = await tailor(TailorRequest(jd_text=JD), profile, llm, FakeEmbeddingProvider())
    assert result.package.llm_calls == 2
    assert result.pre_repair_report is None and result.repaired is False
    assert result.package.usage.input_tokens == 20
    assert result.package.usage.output_tokens == 10


async def test_usage_sums_across_a_repair_round(profile: Profile) -> None:
    llm = FakeLLMProvider([demo_extract(), bad_output(), good_output()])
    result = await tailor(TailorRequest(jd_text=JD), profile, llm, FakeEmbeddingProvider())
    assert result.package.llm_calls == 3
    assert result.package.usage.input_tokens == 30
    assert result.package.usage.output_tokens == 15


async def test_llm_model_is_recorded_on_the_package(profile: Profile) -> None:
    llm = FakeLLMProvider([demo_extract(), good_output()], model="claude-sonnet-5")
    result = await tailor(TailorRequest(jd_text=JD), profile, llm, FakeEmbeddingProvider())
    assert result.package.model == "claude-sonnet-5"


async def test_llm_model_defaults_to_none_without_a_model_attribute(profile: Profile) -> None:
    llm = FakeLLMProvider([demo_extract(), good_output()])
    result = await tailor(TailorRequest(jd_text=JD), profile, llm, FakeEmbeddingProvider())
    assert result.package.model is None


async def test_unrepairable_output_is_blocked(profile: Profile) -> None:
    llm = FakeLLMProvider([demo_extract(), bad_output(), bad_output()])
    result = await tailor(TailorRequest(jd_text=JD), profile, llm, FakeEmbeddingProvider())
    assert result.package.status == "blocked" and result.package.llm_calls == 3
    assert {v.rule for v in result.package.guardrail_report.violations} == {"no-unverified-metrics"}
    assert result.docx == b""  # owner, 2026-09-25: a failing report never persists a DOCX


async def test_budget_exceeded_during_repair_yields_a_blocked_package_not_a_crash(
    profile: Profile,
) -> None:
    """C2: an exhausted budget on the repair path must not raise -- it must return the blocked
    draft, exactly like a MalformedOutputError on the same path already does."""
    llm = FakeLLMProvider([demo_extract(), bad_output(), bad_output()])
    result = await tailor(
        TailorRequest(jd_text=JD),
        profile,
        llm,
        FakeEmbeddingProvider(),
        budget=CallBudget(max_calls=2),
    )
    assert result.package.status == "blocked" and result.package.llm_calls == 2
    assert len(llm.calls) == 2
    assert not result.package.guardrail_report.passed
    # The first report is kept, and the result says no repair ever ran.
    assert result.pre_repair_report is not None and result.repaired is False


async def test_orphan_bullet_blocks_without_docx(profile: Profile) -> None:
    output = good_output()
    output["summary"] = [bullet("Made up.", "ghost").model_dump()]
    llm = FakeLLMProvider([demo_extract(), output, output])
    result = await tailor(TailorRequest(jd_text=JD), profile, llm, FakeEmbeddingProvider())
    assert result.package.status == "blocked" and result.docx == b""
    assert any(v.rule == "provenance" for v in result.package.guardrail_report.violations)


async def test_regeneration_increments_version_and_passes_feedback(profile: Profile) -> None:
    first = await tailor(
        TailorRequest(jd_text=JD),
        profile,
        FakeLLMProvider([demo_extract(), good_output()]),
        FakeEmbeddingProvider(),
    )
    llm = FakeLLMProvider([demo_extract(), good_output()])
    second = await tailor(
        TailorRequest(
            jd_text=JD, feedback="lean harder on migration", previous_package=first.package
        ),
        profile,
        llm,
        FakeEmbeddingProvider(),
    )
    assert second.package.version == 2
    assert "lean harder on migration" in llm.calls[1].messages[0].content
    assert "<previous_resume>" in llm.calls[1].messages[0].content


async def test_unknown_track_raises(profile: Profile) -> None:
    with pytest.raises(ProfileError):
        await tailor(
            TailorRequest(jd_text=JD, track_id="nope"),
            profile,
            FakeLLMProvider([]),
            FakeEmbeddingProvider(),
        )


async def test_cover_note_metric_blocks_even_when_the_resume_is_clean(profile: Profile) -> None:
    output = cover_note_metric_output()
    llm = FakeLLMProvider([demo_extract(), output, output])
    result = await tailor(TailorRequest(jd_text=JD), profile, llm, FakeEmbeddingProvider())
    assert result.package.status == "blocked" and result.package.llm_calls == 3
    violations = result.package.guardrail_report.violations
    assert [v.path for v in violations] == ["cover_note"]
    assert violations[0].rule == "no-unverified-metrics" and violations[0].block_id is None
    assert "37%" in violations[0].message
    assert result.docx == b""  # the resume is clean; only the cover note failed


async def test_missing_role_block_triggers_repair_and_restoring_it_passes(
    profile: Profile,
) -> None:
    llm = FakeLLMProvider([demo_extract(), missing_role_output(), good_output()])
    result = await tailor(TailorRequest(jd_text=JD), profile, llm, FakeEmbeddingProvider())
    assert result.package.status == "draft" and result.package.llm_calls == 3
    assert result.package.guardrail_report.passed
    assert result.repaired is True
    assert result.pre_repair_report is not None
    first = [v for v in result.pre_repair_report.violations if v.rule == "completeness"]
    assert [v.block_id for v in first] == ["acme-data-pm"]
    # The violation itself, not REPAIR_INSTRUCTIONS' own mention of the word, must reach the model.
    sent = llm.calls[2].messages[0].content
    assert "[completeness]" in sent and "was selected but does not appear in Experience" in sent
    assert "acme-data-pm" in sent


async def test_repair_that_drops_a_different_block_is_blocked_with_the_post_repair_report(
    profile: Profile,
) -> None:
    llm = FakeLLMProvider([demo_extract(), missing_role_output(), repair_drops_credential_output()])
    result = await tailor(TailorRequest(jd_text=JD), profile, llm, FakeEmbeddingProvider())
    assert result.package.status == "blocked" and result.package.llm_calls == 3
    violations = result.package.guardrail_report.violations
    assert any(v.rule == "completeness" and v.block_id == "cred-pmp" for v in violations)
    assert not any(v.block_id == "acme-data-pm" for v in violations)  # the role was restored
    assert result.docx == b""  # a failing report never persists a DOCX, whatever rule failed


# --- tune mode -------------------------------------------------------------------------------

CLEAN_BULLET = "Led the Snowflake migration for 12 teams, reducing warehouse cost 30%."
DIRTY_BULLET = "Led the Snowflake migration for 45 teams, reducing warehouse cost 30%."


def _source() -> tuple[SourceDocument, bytes]:
    data = build_fixture_docx()
    return parse_docx(data, "resume.docx"), data


def tune_output(text: str) -> dict[str, Any]:
    return TuneOutput(
        edits=[ProposedEdit(paragraph_id="p9", text=text, reason="mirrors the JD")],
        cover_note="I have led Snowflake migrations end to end for platform teams.",
        change_log="Emphasised the migration.",
        answers=[AnswerItem(key="why_this_company", value="Data.")],
    ).model_dump(mode="json")


def tune_request(doc: SourceDocument, data: bytes, **kwargs: Any) -> TailorRequest:
    return TailorRequest(jd_text=JD, mode="tune", source_document=doc, source_docx=data, **kwargs)


def _docx_paragraph_texts(data: bytes) -> list[str]:
    return [p.text for p in Document(io.BytesIO(data)).paragraphs]


async def test_tune_mode_happy_path_uses_two_calls(profile: Profile) -> None:
    doc, data = _source()
    llm = FakeLLMProvider([demo_extract(), tune_output(CLEAN_BULLET)])
    steps: list[str] = []

    async def on_step(name: str) -> None:
        steps.append(name)

    result = await tailor(
        tune_request(doc, data), profile, llm, FakeEmbeddingProvider(), on_step=on_step
    )
    package = result.package
    assert package.mode == "tune"
    assert package.status == "draft" and package.guardrail_report.passed
    assert package.llm_calls == 2 and len(llm.calls) == 2
    assert package.usage.input_tokens == 20 and package.usage.output_tokens == 10
    assert package.track_id == "data-pm"
    assert [e.paragraph_id for e in package.edits] == ["p9"]
    assert (
        package.edits[0].before
        == "Led the Snowflake migration for 12 teams, cutting warehouse cost 30%."
    )
    assert package.edits[0].after == CLEAN_BULLET
    assert result.edits == package.edits
    assert package.source_document is not None
    assert package.resume.header.name == "Maya Chen"
    assert any(
        b.text == CLEAN_BULLET
        for s in package.resume.sections
        for e in s.entries
        for b in e.bullets
    )
    assert result.docx[:2] == b"PK"
    assert CLEAN_BULLET in _docx_paragraph_texts(result.docx)
    assert steps == ["extract", "tune", "validate", "render"]
    assert result.selection.block_ids == []


async def test_tune_mode_repair_path(profile: Profile) -> None:
    doc, data = _source()
    llm = FakeLLMProvider([demo_extract(), tune_output(DIRTY_BULLET), tune_output(CLEAN_BULLET)])
    result = await tailor(tune_request(doc, data), profile, llm, FakeEmbeddingProvider())
    assert result.package.status == "draft" and result.package.llm_calls == 3
    assert result.package.edits[0].after == CLEAN_BULLET
    repair_call = llm.calls[2]
    assert repair_call.output_schema is TuneOutput
    assert "45" in repair_call.messages[0].content
    assert repair_call.system == llm.calls[1].system  # same cached system blocks as tune


async def test_tune_mode_unrepairable_is_blocked(profile: Profile) -> None:
    doc, data = _source()
    llm = FakeLLMProvider([demo_extract(), tune_output(DIRTY_BULLET), tune_output(DIRTY_BULLET)])
    result = await tailor(tune_request(doc, data), profile, llm, FakeEmbeddingProvider())
    assert result.package.status == "blocked" and result.package.llm_calls == 3
    assert {v.rule for v in result.package.guardrail_report.violations} == {"no-new-numbers"}
    assert result.docx == b""  # nothing safe to write back into the user's document


async def test_tune_mode_malformed_repair_keeps_the_blocked_draft(profile: Profile) -> None:
    doc, data = _source()
    llm = FakeLLMProvider([demo_extract(), tune_output(DIRTY_BULLET), MALFORMED])
    result = await tailor(tune_request(doc, data), profile, llm, FakeEmbeddingProvider())
    assert result.package.status == "blocked" and result.package.llm_calls == 3
    assert not result.package.guardrail_report.passed


async def test_tune_mode_budget_exhausted_before_repair_yields_a_blocked_package(
    profile: Profile,
) -> None:
    doc, data = _source()
    llm = FakeLLMProvider([demo_extract(), tune_output(DIRTY_BULLET), tune_output(CLEAN_BULLET)])
    result = await tailor(
        tune_request(doc, data),
        profile,
        llm,
        FakeEmbeddingProvider(),
        budget=CallBudget(max_calls=2),
    )
    assert result.package.status == "blocked" and result.package.llm_calls == 2
    assert len(llm.calls) == 2 and result.docx == b""


async def test_tune_mode_regeneration_passes_previous_edits(profile: Profile) -> None:
    doc, data = _source()
    first = await tailor(
        tune_request(doc, data),
        profile,
        FakeLLMProvider([demo_extract(), tune_output(CLEAN_BULLET)]),
        FakeEmbeddingProvider(),
    )
    llm = FakeLLMProvider([demo_extract(), tune_output(CLEAN_BULLET)])
    second = await tailor(
        tune_request(doc, data, feedback="shorter", previous_package=first.package),
        profile,
        llm,
        FakeEmbeddingProvider(),
    )
    assert second.package.version == 2
    content = llm.calls[1].messages[0].content
    assert "<previous_edits>" in content and "shorter" in content


def test_tune_mode_requires_document() -> None:
    with pytest.raises(ValidationError):
        TailorRequest(jd_text=JD, mode="tune")
    with pytest.raises(ValidationError):
        TailorRequest(jd_text=JD, mode="tune", source_document=_source()[0])


def test_source_docx_is_excluded_from_json_dumps() -> None:
    doc, data = _source()
    dumped = tune_request(doc, data).model_dump(mode="json")
    assert "source_docx" not in dumped
    assert b"PK" not in tune_request(doc, data).model_dump_json().encode()


async def test_tune_mode_render_failure_yields_an_empty_docx_instead_of_raising(
    profile: Profile,
) -> None:
    """A `source_document` that does not match `source_docx` (stale upload, cached parse) must
    produce a package the human can read, not a failed task -- mirrors blocks mode's
    OrphanBulletError guard."""
    doc, _data = _source()
    short = io.BytesIO()
    stub = Document()
    stub.add_paragraph("MAYA CHEN")
    stub.save(short)
    llm = FakeLLMProvider([demo_extract(), tune_output(CLEAN_BULLET)])
    result = await tailor(
        tune_request(doc, short.getvalue()), profile, llm, FakeEmbeddingProvider()
    )
    assert result.package.guardrail_report.passed  # p9 is a bullet in the parsed document
    assert result.docx == b""
    assert [e.paragraph_id for e in result.package.edits] == ["p9"]
