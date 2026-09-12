import io

import pytest
from docx import Document
from helpers_docx import build_fixture_docx

from rhapto.engine.document import parse_docx
from rhapto.engine.render.tune_docx import render_tuned_docx
from rhapto.engine.types import EngineError
from rhapto.models.source_document import Edit


def test_render_replaces_only_edited_paragraphs_and_keeps_formatting() -> None:
    source = build_fixture_docx()
    doc = parse_docx(source, "resume.docx")
    summary = next(p for p in doc.paragraphs if p.role == "summary")
    title = next(p for p in doc.paragraphs if p.role == "entry_title")
    out = render_tuned_docx(
        source,
        [
            Edit(
                paragraph_id=summary.id,
                before=summary.text,
                after="Data program leader.",
                reason="r",
            )
        ],
    )
    before = [p.text for p in Document(io.BytesIO(source)).paragraphs]
    after_doc = Document(io.BytesIO(out))
    after = [p.text for p in after_doc.paragraphs]
    idx = int(summary.id[1:])
    assert after[idx] == "Data program leader."
    assert [t for i, t in enumerate(after) if i != idx] == [
        t for i, t in enumerate(before) if i != idx
    ]
    title_para = after_doc.paragraphs[int(title.id[1:])]
    assert title_para.runs[0].bold is True  # untouched paragraph formatting intact
    name_run = after_doc.paragraphs[0].runs[0]
    assert name_run.bold is True and name_run.font.size.pt == 17


def test_render_rejects_unknown_paragraph() -> None:
    with pytest.raises(EngineError):
        render_tuned_docx(
            build_fixture_docx(), [Edit(paragraph_id="p999", before="", after="x", reason="r")]
        )
