"""Builds a small fictional resume DOCX in memory, shaped like a real single-column template."""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def build_fixture_docx() -> bytes:
    d = Document()
    p = d.add_paragraph()
    r = p.add_run("MAYA CHEN")
    r.bold = True
    r.font.size = Pt(17)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = d.add_paragraph("Denver, CO  |  555 0100  |  maya.chen@example.com")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for heading, body in (
        ("PROFESSIONAL SUMMARY", None),
        (
            None,
            "Senior Data Program Manager with 8 years leading warehouse migrations and analytics programs.",
        ),
        ("CORE COMPETENCIES", None),
        (None, "Program Delivery:  Integrated Planning  •  RAID Management  •  Cutover"),
        ("PROFESSIONAL EXPERIENCE", None),
        ("TITLE", "Senior Data Program Manager\t2019 – 2025"),
        (None, "Acme Analytics  |  Denver, CO"),
        ("BULLET", "Led the Snowflake migration for 12 teams, cutting warehouse cost 30%."),
        ("BULLET", "Ran the analytics roadmap with finance and engineering stakeholders."),
        ("TECHNICAL SKILLS", None),
        (None, "Platforms:  Snowflake, dbt, Airflow"),
        ("EDUCATION", None),
        (None, "B.S. Computer Science  |  Example University"),
    ):
        if heading == "TITLE":
            t = d.add_paragraph()
            run = t.add_run(body)
            run.bold = True
        elif heading == "BULLET":
            d.add_paragraph(body, style="List Bullet")
        elif heading:
            h = d.add_paragraph()
            run = h.add_run(heading)
            run.bold = True
        else:
            d.add_paragraph(body)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
