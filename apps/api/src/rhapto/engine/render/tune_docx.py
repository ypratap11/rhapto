"""Write tune-mode edits back into the user's DOCX, preserving everything else."""

from __future__ import annotations

import io

from docx import Document
from docx.oxml.ns import qn

from rhapto.engine.types import EngineError
from rhapto.models.source_document import Edit


def _set_text(paragraph, text: str) -> None:  # type: ignore[no-untyped-def]
    runs = list(paragraph.runs)
    for hyperlink in paragraph._p.findall(qn("w:hyperlink")):
        paragraph._p.remove(hyperlink)
    if not runs:
        paragraph.add_run(text)
        return
    for run in runs[1:]:
        run._r.getparent().remove(run._r)
    first = runs[0]
    for child in list(first._r):
        if child.tag != qn("w:rPr"):
            first._r.remove(child)
    for i, chunk in enumerate(text.split("\t")):
        if i:
            first._r.append(first._r.makeelement(qn("w:tab"), {}))
        t = first._r.makeelement(qn("w:t"), {qn("xml:space"): "preserve"})
        t.text = chunk
        first._r.append(t)


def render_tuned_docx(source: bytes, edits: list[Edit]) -> bytes:
    document = Document(io.BytesIO(source))
    paragraphs = document.paragraphs
    for edit in edits:
        try:
            index = int(edit.paragraph_id[1:])
        except ValueError as exc:
            raise EngineError(f"bad paragraph id {edit.paragraph_id!r}") from exc
        if not 0 <= index < len(paragraphs):
            raise EngineError(f"paragraph {edit.paragraph_id} is not in the document")
        _set_text(paragraphs[index], edit.after)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
