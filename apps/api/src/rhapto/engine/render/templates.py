"""DOCX layouts, chosen per resume base via `ResumeBase.style["template"]`.

A template owns presentation only: which headings a section gets, how the name and dates sit on
the page, what is bold. The content it renders -- sections, entries, bullets, provenance -- is the
same `ResumeDocument` whichever template runs, so adding a layout is a rendering function and never
a fork of the writer.

`classic` is the layout Rhapto has always produced and stays the default, so a profile that does
not ask for a template renders exactly as it did before this module existed.
"""

from __future__ import annotations

from collections.abc import Callable, Mapping
from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt

from rhapto.models.resume_document import ResumeDocument, ResumeEntry

#: ATS-safe headings per section kind. The section title the LLM wrote is ignored; `kind` is a
#: validated Literal, so a heading can never be a fabricated label.
CLASSIC_HEADINGS = {
    "experience": "Experience",
    "projects": "Projects",
    "skills": "Skills",
    "credentials": "Credentials",
}

EXECUTIVE_HEADINGS = {
    "experience": "Professional Experience",
    "projects": "Selected Projects",
    "skills": "Areas of Depth",
    "credentials": "Education & Certifications",
}

#: Where the right-aligned date tab sits, in inches from the left margin. Page width less the
#: 0.8in margins on a US Letter page.
_RIGHT_TAB = Inches(6.9)


def _entry_head(entry: ResumeEntry) -> str:
    return " | ".join(filter(None, [entry.role or entry.title, entry.org, entry.period]))


def _bullet(doc: Any, text: str, *, bold_lead: bool) -> None:
    """One bullet. With `bold_lead`, the opening label sentence is bolded.

    The composer is asked to open a bullet with a short label ("Executive decisions. Drove the
    go/no-go..."), which reads well in any layout; only this template makes it visually distinct.
    A bullet without a label renders plain rather than bolding an arbitrary prefix.
    """
    paragraph = doc.add_paragraph(style="List Bullet")
    head, sep, rest = text.partition(". ")
    if bold_lead and sep and 0 < len(head) <= 60:
        paragraph.add_run(head + ".").bold = True
        paragraph.add_run(" " + rest)
    else:
        paragraph.add_run(text)


def render_classic(doc: Any, resume: ResumeDocument) -> None:
    """The original layout: left-aligned name, uppercase headings, one combined entry line."""
    name = doc.add_paragraph()
    name_run = name.add_run(resume.header.name)
    name_run.bold = True
    name_run.font.size = Pt(16)

    header = resume.header
    contact = " | ".join(filter(None, [header.email, header.phone, header.location, *header.links]))
    if contact:
        doc.add_paragraph(contact)

    if resume.summary:
        _classic_heading(doc, "Summary")
        doc.add_paragraph(" ".join(b.text for b in resume.summary))

    for section in resume.sections:
        _classic_heading(doc, CLASSIC_HEADINGS.get(section.kind, section.kind))
        for entry in section.entries:
            head = _entry_head(entry)
            if head:
                paragraph = doc.add_paragraph()
                paragraph.add_run(head).bold = True
            for bullet in entry.bullets:
                _bullet(doc, bullet.text, bold_lead=False)


def _classic_heading(doc: Any, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(12)
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(2)


def _executive_heading(doc: Any, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(3)


def render_executive(doc: Any, resume: ResumeDocument) -> None:
    """Centred name, Title Case headings, dates right-tabbed, bold bullet lead-ins."""
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name.add_run(resume.header.name.upper())
    name_run.bold = True
    name_run.font.size = Pt(16)

    header = resume.header
    contact = "  |  ".join(
        filter(None, [header.location, header.phone, header.email, *header.links])
    )
    if contact:
        line = doc.add_paragraph(contact)
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if resume.summary:
        _executive_heading(doc, "Summary")
        doc.add_paragraph(" ".join(b.text for b in resume.summary))

    for section in resume.sections:
        _executive_heading(doc, EXECUTIVE_HEADINGS.get(section.kind, section.kind))
        for entry in section.entries:
            title = entry.role or entry.title
            if title:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.tab_stops.add_tab_stop(
                    _RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT
                )
                paragraph.add_run(title).bold = True
                if entry.period:
                    paragraph.add_run("\t" + entry.period)
            if entry.org:
                doc.add_paragraph(entry.org)
            for bullet in entry.bullets:
                _bullet(doc, bullet.text, bold_lead=True)


Template = Callable[[Any, ResumeDocument], None]

TEMPLATES: Mapping[str, Template] = {
    "classic": render_classic,
    "executive": render_executive,
}

DEFAULT_TEMPLATE = "classic"


def template_for(style: Mapping[str, Any] | None) -> Template:
    """The template a base's style asks for, falling back to the default.

    An unknown name falls back rather than raising: a typo in a profile must not fail a tailoring
    run that has already spent its LLM calls and passed every guardrail.
    """
    name = (style or {}).get("template")
    if not isinstance(name, str):
        return TEMPLATES[DEFAULT_TEMPLATE]
    return TEMPLATES.get(name, TEMPLATES[DEFAULT_TEMPLATE])
